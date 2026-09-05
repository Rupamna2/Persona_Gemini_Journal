"""Tests for Excel Export Streaming Route (Unit 14)."""

import io
from datetime import datetime, timezone, timedelta
from unittest.mock import patch, MagicMock
from fastapi.testclient import TestClient
import openpyxl

from backend.main import app
from backend.auth import verify_token
from backend.routes.export import HEADER_COLUMNS

client = TestClient(app)

TEST_UID = "test-export-user-999"


def override_verify_token_valid():
    return {
        "uid": TEST_UID,
        "email": "exportuser@example.com",
        "name": "Export User",
    }


def test_export_unauthorized_returns_401():
    """Unauthenticated GET /api/export returns 401."""
    res = client.get("/api/export")
    assert res.status_code == 401


def test_export_invalid_range_returns_400():
    """Invalid range parameter returns 400."""
    app.dependency_overrides[verify_token] = override_verify_token_valid
    try:
        res = client.get("/api/export?range=yearly")
        assert res.status_code == 400
        assert "Invalid range parameter" in res.json()["detail"]
    finally:
        app.dependency_overrides.clear()


def test_export_generates_valid_excel_workbook():
    """Valid export generates in-memory .xlsx with headers and rows."""
    app.dependency_overrides[verify_token] = override_verify_token_valid

    now_iso = datetime.now(timezone.utc).isoformat()
    yesterday_iso = (datetime.now(timezone.utc) - timedelta(days=1)).isoformat()

    mock_doc1 = MagicMock()
    mock_doc1.id = "j-exp-1"
    mock_doc1.to_dict.return_value = {
        "journalId": "j-exp-1",
        "date": "2026-08-30",
        "mode": "DecisionMaking",
        "title": "Scaling Cloud Run Infrastructure",
        "summary": {
            "mood_score": 8.5,
            "topic": "DevOps & Cloud Run",
            "key_insights": ["Container concurrency tuning reduced p99 latency"],
            "action_items": ["Set min-instances to 1 for zero cold-starts"],
        },
        "createdAt": now_iso,
    }

    mock_doc2 = MagicMock()
    mock_doc2.id = "j-exp-2"
    mock_doc2.to_dict.return_value = {
        "journalId": "j-exp-2",
        "date": "2026-08-29",
        "mode": "Gratitude",
        "title": "Morning Walk & Coffee",
        "summary": {
            "mood_score": 9.0,
            "topic": "Wellbeing",
            "key_insights": ["Consistent sleep routine improved daytime focus"],
            "action_items": [],
        },
        "createdAt": yesterday_iso,
    }

    mock_db = MagicMock()
    mock_query = MagicMock()
    mock_query.stream.return_value = [mock_doc1, mock_doc2]

    mock_journals_coll = MagicMock()
    mock_journals_coll.order_by.return_value = mock_query

    mock_user_doc = MagicMock()
    mock_user_doc.collection.return_value = mock_journals_coll
    mock_db.collection.return_value.document.return_value = mock_user_doc

    with patch("backend.routes.export.get_firestore_client", return_value=mock_db):
        try:
            res = client.get("/api/export?range=weekly")
            assert res.status_code == 200
            assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            assert "attachment; filename=" in res.headers["content-disposition"]

            # Load streamed bytes into openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(res.content))
            ws = wb.active
            assert ws.title == "Personal Reflections"

            # Check header row
            header_vals = [ws.cell(row=1, column=c).value for c in range(1, len(HEADER_COLUMNS) + 1)]
            assert header_vals == HEADER_COLUMNS

            # Check data rows
            assert ws.max_row == 3  # 1 header + 2 data rows
            row1_title = ws.cell(row=2, column=4).value
            assert row1_title == "Scaling Cloud Run Infrastructure"
            row1_score = ws.cell(row=2, column=5).value
            assert float(row1_score) == 8.5

            row2_title = ws.cell(row=3, column=4).value
            assert row2_title == "Morning Walk & Coffee"

            # Check strictly user scoped query
            mock_db.collection.assert_called_with("users")
            mock_db.collection("users").document.assert_called_with(TEST_UID)
            mock_user_doc.collection.assert_called_with("journals")

        finally:
            app.dependency_overrides.clear()


def test_export_empty_journals_returns_valid_header_workbook():
    """When user has 0 journals, export returns a valid header-only workbook (not an error)."""
    app.dependency_overrides[verify_token] = override_verify_token_valid

    mock_db = MagicMock()
    mock_query = MagicMock()
    mock_query.stream.return_value = []

    mock_journals_coll = MagicMock()
    mock_journals_coll.order_by.return_value = mock_query

    mock_user_doc = MagicMock()
    mock_user_doc.collection.return_value = mock_journals_coll
    mock_db.collection.return_value.document.return_value = mock_user_doc

    with patch("backend.routes.export.get_firestore_client", return_value=mock_db):
        try:
            res = client.get("/api/export?range=all")
            assert res.status_code == 200

            wb = openpyxl.load_workbook(io.BytesIO(res.content))
            ws = wb.active
            assert ws.max_row == 1  # only header
            header_vals = [ws.cell(row=1, column=c).value for c in range(1, len(HEADER_COLUMNS) + 1)]
            assert header_vals == HEADER_COLUMNS

        finally:
            app.dependency_overrides.clear()
